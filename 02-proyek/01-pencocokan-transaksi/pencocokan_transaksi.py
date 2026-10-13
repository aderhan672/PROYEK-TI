#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Aplikasi Desktop: Pencocokan Transaksi v4.3.2
Reference-Aware + Multi-Pass + DFS/MILP Hybrid + Deep Search
"""
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import csv, os, re, sqlite3, threading, queue, time, io, contextlib
from datetime import datetime
from collections import defaultdict
from itertools import combinations

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from scipy.optimize import milp, LinearConstraint, Bounds
    import numpy as np
    HAS_SCIPY = True
except ImportError:
    HAS_SCIPY = False

DB_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data.db")


# =========================================================
# UTILITAS
# =========================================================
def parse_number(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s:
        return None
    s = re.sub(r'(?i)\brp\.?\b', '', s)
    s = re.sub(r'(?i)\bidr\b', '', s)
    s = s.replace(' ', '').replace('\u00a0', '')
    neg = False
    if s.startswith('(') and s.endswith(')'):
        neg = True; s = s[1:-1]
    if s.startswith('-'):
        neg = True; s = s[1:]
    if ',' in s and '.' in s:
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    elif ',' in s and '.' not in s:
        parts = s.split(',')
        if len(parts) == 2 and len(parts[1]) <= 2:
            s = s.replace(',', '.')
        else:
            s = s.replace(',', '')
    elif '.' in s and ',' not in s:
        parts = s.split('.')
        if len(parts) > 1 and all(len(p) == 3 for p in parts[1:] if p):
            s = s.replace('.', '')
    s = re.sub(r'[^0-9.]', '', s)
    if s.count('.') > 1:
        head, _, tail = s.rpartition('.')
        s = head.replace('.', '') + '.' + tail
    try:
        val = float(s)
        return -val if neg else val
    except ValueError:
        return None


def norm_supplier(s):
    return re.sub(r'\s+', ' ', str(s or '').strip().upper())


def to_cents(v):
    return 0 if v is None else int(round(v * 100))


def fmt_money(v):
    if v is None:
        return ''
    try:
        return f"{v:,.2f}"
    except (TypeError, ValueError):
        return str(v)


REF_PATTERNS = [
    re.compile(r'YMID\d+', re.I),
    re.compile(r'\b[A-Z]{2,}[/-]?\d{2,}[/-]?\d+\b', re.I),
    re.compile(r'\b\d{8,}\b'),
    re.compile(r'\b[A-Z]{2,}\d{4,}\b', re.I),
]


def extract_references(keterangan):
    if not keterangan:
        return set()
    s = str(keterangan).upper()
    refs = set()
    for pat in REF_PATTERNS:
        for m in pat.finditer(s):
            refs.add(m.group().upper())
    return refs


# =========================================================
# DATABASE
# =========================================================
def db_connect():
    conn = sqlite3.connect(DB_FILE, timeout=30.0)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode = WAL")
    conn.execute("PRAGMA synchronous = NORMAL")
    return conn


def db_init():
    conn = db_connect()
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS transaksi (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            sesi_import TEXT, source_file TEXT, sheet_name TEXT, row_excel INTEGER,
            tanggal TEXT, sumber TEXT, supplier TEXT, no_sumber TEXT,
            keterangan TEXT, debit REAL, kredit REAL, balance TEXT
        );
        CREATE INDEX IF NOT EXISTS idx_trx_supplier ON transaksi(supplier);
        CREATE INDEX IF NOT EXISTS idx_trx_sesi ON transaksi(sesi_import);
        CREATE INDEX IF NOT EXISTS idx_trx_row ON transaksi(row_excel);
        CREATE TABLE IF NOT EXISTS hasil_match (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            sesi_match TEXT, supplier TEXT, tipe_group TEXT,
            debit_ids TEXT, credit_ids TEXT, total_nilai REAL,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        );
        CREATE INDEX IF NOT EXISTS idx_match_sesi ON hasil_match(sesi_match);
        CREATE TABLE IF NOT EXISTS tanpa_pasangan (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            sesi_match TEXT, supplier TEXT, transaksi_id INTEGER,
            row_excel INTEGER, nilai REAL, sisi TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        );
        CREATE INDEX IF NOT EXISTS idx_unmatch_sesi ON tanpa_pasangan(sesi_match);
    """)
    conn.commit(); conn.close()


def db_get_suppliers():
    conn = db_connect()
    rows = [r[0] for r in conn.execute(
        "SELECT DISTINCT supplier FROM transaksi WHERE supplier != '' ORDER BY supplier")]
    conn.close(); return rows


def db_count_rows(sesi=None):
    conn = db_connect()
    if sesi:
        n = conn.execute("SELECT COUNT(*) FROM transaksi WHERE sesi_import=?", (sesi,)).fetchone()[0]
    else:
        n = conn.execute("SELECT COUNT(*) FROM transaksi").fetchone()[0]
    conn.close(); return n


def db_get_import_sessions():
    conn = db_connect()
    rows = [dict(r) for r in conn.execute(
        """SELECT sesi_import, source_file, COUNT(*) as jml, MIN(sheet_name) as sheet
           FROM transaksi GROUP BY sesi_import ORDER BY sesi_import DESC""")]
    conn.close(); return rows


def db_clear_all():
    conn = db_connect()
    conn.executescript("""
        DELETE FROM transaksi; DELETE FROM hasil_match;
        DELETE FROM tanpa_pasangan; VACUUM;
    """)
    conn.commit(); conn.close()


def db_get_data(supplier=None, limit=None):
    conn = db_connect()
    sql = "SELECT * FROM transaksi"
    params = []
    if supplier and supplier != 'ALL':
        sql += " WHERE supplier = ?"; params.append(supplier)
    sql += " ORDER BY row_excel"
    if limit:
        sql += f" LIMIT {int(limit)}"
    rows = [dict(r) for r in conn.execute(sql, params)]
    conn.close(); return rows


# =========================================================
# IMPORT EXCEL
# =========================================================
def read_excel_rows(path, sheet_name):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close(); return rows


def db_import_excel(path, sheet_name, header_row, col_map, label,
                    replace=False, progress_cb=None, cancel_event=None):
    rows = read_excel_rows(path, sheet_name)
    if not rows:
        raise ValueError("File kosong.")
    hr = max(1, int(header_row))
    headers = [str(h).strip() if h is not None and str(h).strip() else f"Kolom{i+1}"
               for i, h in enumerate(rows[hr-1])]
    idx = {}
    for key, col in col_map.items():
        idx[key] = headers.index(col) if col and col in headers else None

    def find_col(*keys):
        for k in keys:
            for i, h in enumerate(headers):
                if k.lower() in h.lower():
                    return i
        return None

    idx['tanggal'] = find_col('tanggal', 'date')
    idx['sumber'] = find_col('sumber')
    idx['no_sumber'] = find_col('no. sumber', 'no sumber', 'nomor')
    idx['keterangan'] = find_col('keterangan', 'desc')
    idx['balance'] = find_col('balance')

    data_rows = rows[hr:]
    total = len(data_rows)
    conn = db_connect(); cur = conn.cursor()
    if replace:
        cur.execute("DELETE FROM transaksi")
    sesi = datetime.now().strftime("%Y%m%d_%H%M%S")
    batch = []; inserted = 0; skipped = 0

    for i, row in enumerate(data_rows):
        if cancel_event and cancel_event.is_set():
            conn.rollback(); conn.close()
            raise InterruptedError("Dibatalkan")

        def g(k):
            ix = idx.get(k)
            if ix is None: return None
            return row[ix] if 0 <= ix < len(row) else None

        supplier = norm_supplier(g('supplier'))
        debit = parse_number(g('debit'))
        credit = parse_number(g('kredit'))
        if supplier == '' and (debit is None or debit == 0) and (credit is None or credit == 0):
            skipped += 1; continue
        row_excel = hr + i + 1
        batch.append((sesi, label, sheet_name, row_excel,
                      str(g('tanggal') or ''), str(g('sumber') or ''), supplier,
                      str(g('no_sumber') or ''), str(g('keterangan') or ''),
                      debit if debit is not None else 0.0,
                      credit if credit is not None else 0.0,
                      str(g('balance') or '')))
        inserted += 1
        if len(batch) >= 500:
            cur.executemany("""INSERT INTO transaksi
                (sesi_import,source_file,sheet_name,row_excel,tanggal,sumber,supplier,
                 no_sumber,keterangan,debit,kredit,balance)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""", batch)
            batch = []
            if progress_cb:
                progress_cb(int(100*(i+1)/total), f"Import: {i+1}/{total}")
    if batch:
        cur.executemany("""INSERT INTO transaksi
            (sesi_import,source_file,sheet_name,row_excel,tanggal,sumber,supplier,
             no_sumber,keterangan,debit,kredit,balance)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""", batch)
    conn.commit(); conn.close()
    if progress_cb: progress_cb(100, "Selesai")
    return {'sesi': sesi, 'inserted': inserted, 'skipped': skipped,
            'total_rows': total, 'headers': headers}


# =========================================================
# ALGORITMA v4.3.2 — DFS + MILP HYBRID + DEEP SEARCH
# =========================================================
def _find_subset_dfs(items, target, max_size, min_size=2,
                     time_budget_sec=0.3, cancel_event=None):
    """DFS item-by-item. Cepat untuk kombinasi kecil."""
    if target <= 0 or max_size < min_size:
        return None
    vals = [(it['cents'], it) for it in items if it['cents'] > 0]
    if len(vals) < min_size:
        return None
    vals.sort(key=lambda x: -x[0])
    n = len(vals)
    cents = [v for v, _ in vals]
    items_arr = [it for _, it in vals]

    suffix = [0] * (n + 1)
    for i in range(n - 1, -1, -1):
        suffix[i] = suffix[i + 1] + cents[i]
    if suffix[0] < target:
        return None

    smallest = cents[-1]
    largest = cents[0]
    min_real = max(min_size, (target + largest - 1) // largest)
    max_real = min(max_size, target // smallest if smallest > 0 else max_size)
    if min_real > max_real:
        return None

    deadline = time.time() + time_budget_sec
    nodes = [0]
    picked = []

    def bt(start, remaining, max_sz, dl):
        nodes[0] += 1
        if nodes[0] % 200 == 0:
            if cancel_event and cancel_event.is_set():
                raise InterruptedError("Dibatalkan")
            if time.time() > dl:
                raise TimeoutError()
        if remaining == 0:
            return list(picked) if len(picked) >= min_size else None
        if start >= n or len(picked) >= max_sz:
            return None
        if suffix[start] < remaining:
            return None
        slots = max_sz - len(picked)
        if slots <= 0:
            return None
        if cents[start] * slots < remaining:
            return None
        if smallest * slots > remaining:
            return None
        for i in range(start, n):
            v = cents[i]
            if v > remaining:
                continue
            picked.append(i)
            try:
                res = bt(i + 1, remaining - v, max_sz, dl)
            except (TimeoutError, InterruptedError):
                picked.pop(); raise
            picked.pop()
            if res is not None:
                return res
        return None

    sizes = list(range(min_real, max_real + 1))
    sizes.sort(key=lambda s: (s if s <= 6 else 6 + (s - 6) * 0.1))
    per_size = time_budget_sec / max(1, len(sizes))
    for size in sizes:
        if time.time() > deadline:
            break
        picked.clear()
        try:
            res = bt(0, target, size, time.time() + per_size)
        except TimeoutError:
            continue
        except InterruptedError:
            raise
        if res is not None:
            return [items_arr[i] for i in res]
    return None


def _find_subset_milp(items, target, max_size, min_size=2, time_budget=5.0):
    """
    MILP solver (butuh scipy). Output HiGHS di-suppress.
    v4.3.2: timeout diperpanjang maks 60 detik.
    """
    if not HAS_SCIPY:
        return None
    if target <= 0 or max_size < min_size:
        return None
    vals = [(it['cents'], it) for it in items if it['cents'] > 0]
    if len(vals) < min_size:
        return None
    n = len(vals)
    cents = np.array([v for v, _ in vals], dtype=np.float64)
    items_arr = [it for _, it in vals]

    c = np.zeros(n)
    A = np.vstack([cents, np.ones(n)])
    lb = np.array([float(target), float(min_size)])
    ub = np.array([float(target), float(max_size)])
    constraints = LinearConstraint(A, lb, ub)
    bounds = Bounds(0, 1)
    integrality = np.ones(n)

    # v4.3.2: timeout maks 60 detik (sebelumnya 3 detik)
    milp_timeout = min(60.0, max(5.0, time_budget))

    try:
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(buf):
            res = milp(c=c, constraints=constraints, bounds=bounds,
                       integrality=integrality,
                       options={'time_limit': milp_timeout,
                                'presolve': True,
                                'disp': False})
    except Exception:
        return None
    if res is not None and res.success and res.x is not None:
        picked = [items_arr[i] for i in range(n) if res.x[i] > 0.5]
        if sum(p['cents'] for p in picked) == target:
            return picked
    return None


def find_subset(items, target, max_size, min_size=2,
                time_budget_sec=0.3, cancel_event=None):
    """
    v4.3.2: DFS cepat dulu, fallback MILP jika gagal.
    MILP timeout diperpanjang sampai 60 detik.
    """
    small_max = min(max_size, 6)
    result = _find_subset_dfs(items, target, small_max, min_size,
                              min(time_budget_sec, 0.2), cancel_event)
    if result is not None:
        return result

    if max_size > 6:
        remaining = max(0.1, time_budget_sec - 0.2)
        result = _find_subset_dfs(items, target, max_size, min_size,
                                  remaining, cancel_event)
        if result is not None:
            return result

    if HAS_SCIPY:
        milp_budget = min(60.0, max(5.0, time_budget_sec * 5))
        return _find_subset_milp(items, target, max_size, min_size,
                                 time_budget=milp_budget)
    return None


def _match_side(dlist, clist, supp, max_group, time_budget,
                cancel_event, groups, used_d, used_c, progress_tag="",
                progress_cb=None, base_pct=0.0, span_pct=100.0,
                enable_deep=False):
    def check_cancel():
        if cancel_event and cancel_event.is_set():
            raise InterruptedError("Dibatalkan")

    def report(sub_pct, msg):
        if progress_cb:
            pct = int(base_pct + span_pct * sub_pct / 100.0)
            progress_cb(pct, msg)

    # PASS 1: 1↔1
    c_by_cents = defaultdict(list)
    for c in clist:
        c_by_cents[c['cents']].append(c)
    for d in dlist:
        if d['id'] in used_d: continue
        bucket = c_by_cents.get(d['cents'])
        if not bucket: continue
        for c in bucket:
            if c['id'] not in used_c:
                used_d.add(d['id']); used_c.add(c['id'])
                groups.append({'supplier': supp, 'debit_ids': [d['id']],
                               'credit_ids': [c['id']], 'total': d['cents']/100.0,
                               'size': '1↔1'})
                break

    # PASS 2: 1↔N kecil (N=2..4)
    if max_group >= 2:
        unused_d = [d for d in dlist if d['id'] not in used_d]
        unused_d.sort(key=lambda x: -x['cents'])
        total_d = len(unused_d) or 1
        for idx_d, d in enumerate(unused_d):
            check_cancel()
            if d['id'] in used_d: continue
            if d['cents'] <= 0: continue
            avail_c = [c for c in clist if c['id'] not in used_c]
            if len(avail_c) < 2: break
            n_max = min(4, max_group)
            if idx_d % 5 == 0:
                report(int(5 * idx_d / total_d),
                       f"{progress_tag} — 1↔N kecil {idx_d+1}/{total_d}")
            try:
                result = find_subset(avail_c, d['cents'], n_max, min_size=2,
                                     time_budget_sec=min(time_budget, 0.15),
                                     cancel_event=cancel_event)
            except InterruptedError:
                raise
            if result:
                used_d.add(d['id'])
                for c in result:
                    used_c.add(c['id'])
                groups.append({'supplier': supp, 'debit_ids': [d['id']],
                               'credit_ids': [c['id'] for c in result],
                               'total': d['cents']/100.0,
                               'size': f'1↔{len(result)}'})

    # PASS 3: M↔1 kecil (M=2..4)
    if max_group >= 2:
        unused_c = [c for c in clist if c['id'] not in used_c]
        unused_c.sort(key=lambda x: -x['cents'])
        total_c = len(unused_c) or 1
        for idx_c, c in enumerate(unused_c):
            check_cancel()
            if c['id'] in used_c: continue
            if c['cents'] <= 0: continue
            avail_d = [d for d in dlist if d['id'] not in used_d]
            if len(avail_d) < 2: break
            n_max = min(4, max_group)
            if idx_c % 5 == 0:
                report(5 + int(5 * idx_c / total_c),
                       f"{progress_tag} — M↔1 kecil {idx_c+1}/{total_c}")
            try:
                result = find_subset(avail_d, c['cents'], n_max, min_size=2,
                                     time_budget_sec=min(time_budget, 0.15),
                                     cancel_event=cancel_event)
            except InterruptedError:
                raise
            if result:
                used_c.add(c['id'])
                for d in result:
                    used_d.add(d['id'])
                groups.append({'supplier': supp,
                               'debit_ids': [d['id'] for d in result],
                               'credit_ids': [c['id']],
                               'total': c['cents']/100.0,
                               'size': f'{len(result)}↔1'})

    # PASS 4: M↔N kecil
    if max_group >= 2:
        report(10, f"{progress_tag} — M↔N kecil...")
        unused_d = [d for d in dlist if d['id'] not in used_d]
        unused_c = [c for c in clist if c['id'] not in used_c]
        MAX_DSIDE = min(3, max_group)
        MAX_CSIDE = min(3, max_group)

        d_cache = defaultdict(list)
        for k in range(2, MAX_DSIDE + 1):
            for combo in combinations(unused_d, k):
                total = sum(x['cents'] for x in combo)
                d_cache[total].append(combo)

        c_cache = defaultdict(list)
        for k in range(2, MAX_CSIDE + 1):
            for combo in combinations(unused_c, k):
                total = sum(x['cents'] for x in combo)
                c_cache[total].append(combo)

        for total in sorted(d_cache.keys()):
            check_cancel()
            if total not in c_cache: continue
            for d_combo in d_cache[total]:
                if any(x['id'] in used_d for x in d_combo): continue
                matched_c = None
                for c_combo in c_cache[total]:
                    if any(x['id'] in used_c for x in c_combo): continue
                    matched_c = c_combo
                    break
                if matched_c is None: continue
                for d in d_combo: used_d.add(d['id'])
                for c in matched_c: used_c.add(c['id'])
                groups.append({'supplier': supp,
                               'debit_ids': [d['id'] for d in d_combo],
                               'credit_ids': [c['id'] for c in matched_c],
                               'total': total/100.0,
                               'size': f'{len(d_combo)}↔{len(matched_c)}'})

    # PASS 5: 1↔N besar
    if max_group >= 5:
        unused_d = [d for d in dlist if d['id'] not in used_d]
        unused_d.sort(key=lambda x: -x['cents'])
        total_d = len(unused_d) or 1
        for idx_d, d in enumerate(unused_d):
            check_cancel()
            if d['id'] in used_d: continue
            if d['cents'] <= 0: continue
            avail_c = [c for c in clist if c['id'] not in used_c]
            if len(avail_c) < 3: break
            report(10 + int(40 * idx_d / total_d),
                   f"{progress_tag} — 1↔N besar {idx_d+1}/{total_d} "
                   f"(Rp {d['cents']/100:,.0f})")
            try:
                result = find_subset(avail_c, d['cents'], max_group, min_size=4,
                                     time_budget_sec=time_budget,
                                     cancel_event=cancel_event)
            except InterruptedError:
                raise
            if result:
                used_d.add(d['id'])
                for c in result:
                    used_c.add(c['id'])
                groups.append({'supplier': supp, 'debit_ids': [d['id']],
                               'credit_ids': [c['id'] for c in result],
                               'total': d['cents']/100.0,
                               'size': f'1↔{len(result)}'})

    # PASS 6: M↔1 besar
    if max_group >= 5:
        unused_c = [c for c in clist if c['id'] not in used_c]
        unused_c.sort(key=lambda x: -x['cents'])
        total_c = len(unused_c) or 1
        for idx_c, c in enumerate(unused_c):
            check_cancel()
            if c['id'] in used_c: continue
            if c['cents'] <= 0: continue
            avail_d = [d for d in dlist if d['id'] not in used_d]
            if len(avail_d) < 3: break
            report(50 + int(40 * idx_c / total_c),
                   f"{progress_tag} — M↔1 besar {idx_c+1}/{total_c} "
                   f"(Rp {c['cents']/100:,.0f})")
            try:
                result = find_subset(avail_d, c['cents'], max_group, min_size=4,
                                     time_budget_sec=time_budget,
                                     cancel_event=cancel_event)
            except InterruptedError:
                raise
            if result:
                used_c.add(c['id'])
                for d in result:
                    used_d.add(d['id'])
                groups.append({'supplier': supp,
                               'debit_ids': [d['id'] for d in result],
                               'credit_ids': [c['id']],
                               'total': c['cents']/100.0,
                               'size': f'{len(result)}↔1'})

    # PASS 7: DEEP SEARCH — MILP 30 detik untuk sisa debit yang belum match
    if enable_deep and HAS_SCIPY and max_group >= 5:
        unused_d = [d for d in dlist if d['id'] not in used_d]
        unused_d.sort(key=lambda x: -x['cents'])
        total_d = len(unused_d) or 1
        for idx_d, d in enumerate(unused_d):
            check_cancel()
            if d['id'] in used_d: continue
            if d['cents'] <= 0: continue
            avail_c = [c for c in clist if c['id'] not in used_c]
            if len(avail_c) < 2: break
            report(90 + int(9 * idx_d / total_d),
                   f"{progress_tag} — DEEP {idx_d+1}/{total_d} "
                   f"(Rp {d['cents']/100:,.0f}, {len(avail_c)} kredit)")
            try:
                result = _find_subset_milp(avail_c, d['cents'], max_group,
                                            min_size=2, time_budget=30.0)
            except InterruptedError:
                raise
            if result:
                used_d.add(d['id'])
                for c in result:
                    used_c.add(c['id'])
                groups.append({'supplier': supp, 'debit_ids': [d['id']],
                               'credit_ids': [c['id'] for c in result],
                               'total': d['cents']/100.0,
                               'size': f'1↔{len(result)}'})

    report(100, f"{progress_tag} — selesai")


def match_rows(rows, max_group=5, time_budget=0.2,
               cancel_event=None, progress_cb=None):
    def check_cancel():
        if cancel_event and cancel_event.is_set():
            raise InterruptedError("Dibatalkan")

    by_supplier = defaultdict(lambda: {'d': [], 'c': []})
    for r in rows:
        d = r['debit'] or 0
        c = r['kredit'] or 0
        if d > 0.005 and c <= 0.005:
            by_supplier[r['supplier']]['d'].append({
                'id': r['id'], 'row': r['row_excel'], 'cents': to_cents(d),
                'raw': r, 'refs': extract_references(r.get('keterangan'))
            })
        elif c > 0.005 and d <= 0.005:
            by_supplier[r['supplier']]['c'].append({
                'id': r['id'], 'row': r['row_excel'], 'cents': to_cents(c),
                'raw': r, 'refs': extract_references(r.get('keterangan'))
            })

    suppliers = sorted(by_supplier.keys())
    total_supp = len(suppliers) or 1
    groups = []
    used_d = set()
    used_c = set()

    for s_idx, supp in enumerate(suppliers):
        check_cancel()
        dlist_all = by_supplier[supp]['d']
        clist_all = by_supplier[supp]['c']

        base = 100.0 * s_idx / total_supp
        span = 100.0 / total_supp

        if progress_cb:
            progress_cb(int(base),
                f"[{s_idx+1}/{total_supp}] {supp[:30]} — "
                f"{len(dlist_all)}D/{len(clist_all)}K")

        # SUB-PASS A: by reference
        ref_d = defaultdict(list)
        ref_c = defaultdict(list)
        for d in dlist_all:
            for rf in d['refs']:
                ref_d[rf].append(d)
        for c in clist_all:
            for rf in c['refs']:
                ref_c[rf].append(c)

        common_refs = sorted(set(ref_d.keys()) & set(ref_c.keys()))
        for rf in common_refs:
            check_cancel()
            sub_d = [d for d in ref_d[rf] if d['id'] not in used_d]
            sub_c = [c for c in ref_c[rf] if c['id'] not in used_c]
            if not sub_d or not sub_c:
                continue
            _match_side(sub_d, sub_c, supp, max_group, time_budget,
                        cancel_event, groups, used_d, used_c, f"ref:{rf}",
                        progress_cb=progress_cb,
                        base_pct=base, span_pct=span * 0.30,
                        enable_deep=False)

        # SUB-PASS B: level supplier + DEEP SEARCH
        rem_d = [d for d in dlist_all if d['id'] not in used_d]
        rem_c = [c for c in clist_all if c['id'] not in used_c]
        if rem_d and rem_c:
            _match_side(rem_d, rem_c, supp, max_group, time_budget,
                        cancel_event, groups, used_d, used_c, "supplier",
                        progress_cb=progress_cb,
                        base_pct=base + span * 0.30, span_pct=span * 0.70,
                        enable_deep=True)

        if progress_cb:
            progress_cb(int(base + span),
                f"[{s_idx+1}/{total_supp}] {supp[:30]} — {len(groups)} grup")

    unmatched_d = [d for grp in by_supplier.values() for d in grp['d']
                   if d['id'] not in used_d]
    unmatched_c = [c for grp in by_supplier.values() for c in grp['c']
                   if c['id'] not in used_c]
    return groups, unmatched_d, unmatched_c


# =========================================================
# GUI
# =========================================================
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        scipy_status = "MILP aktif" if HAS_SCIPY else "MILP off (install scipy)"
        self.title(f"Pencocokan Transaksi v4.3.2 — DFS+MILP+Deep ({scipy_status})")
        self.geometry("1280x760"); self.minsize(1024, 600)
        self.cancel_event = threading.Event()
        self.worker_thread = None
        self.ui_queue = queue.Queue()
        self.imported_file = None
        self.imported_sheet = None
        self.last_match_sesi = None
        db_init()
        self._build_ui()
        self._refresh_suppliers()
        self._poll_queue()
        self._refresh_data_tab()

    def _build_ui(self):
        menubar = tk.Menu(self)
        fm = tk.Menu(menubar, tearoff=0)
        fm.add_command(label="📥  Import Excel ke Database…", command=self._show_import_dialog)
        fm.add_command(label="🗑  Hapus Semua Data…", command=self._clear_all_data)
        fm.add_separator()
        fm.add_command(label="❌  Keluar", command=self.destroy)
        menubar.add_cascade(label="File", menu=fm)
        hm = tk.Menu(menubar, tearoff=0)
        hm.add_command(label="Tentang", command=self._show_about)
        menubar.add_cascade(label="Bantuan", menu=hm)
        self.config(menu=menubar)

        self.status_var = tk.StringVar(value="Siap.")
        ttk.Label(self, textvariable=self.status_var, relief='sunken',
                  anchor='w', padding=(6, 2)).pack(fill='x', side='bottom')

        self.nb = ttk.Notebook(self); self.nb.pack(fill='both', expand=True)
        self.nb.bind('<<NotebookTabChanged>>', self._on_tab_changed)
        self.tab_data = ttk.Frame(self.nb, padding=8)
        self.tab_match = ttk.Frame(self.nb, padding=8)
        self.tab_result = ttk.Frame(self.nb, padding=8)
        self.nb.add(self.tab_data, text="  📊 Data  ")
        self.nb.add(self.tab_match, text="  🔗 Pencocokan  ")
        self.nb.add(self.tab_result, text="  📋 Hasil  ")
        self._build_tab_data(); self._build_tab_match(); self._build_tab_result()

    def _on_tab_changed(self, e):
        try:
            i = self.nb.index(self.nb.select())
            if i == 0: self._refresh_data_tab()
            elif i == 1: self._refresh_suppliers()
        except Exception: pass

    def _refresh_suppliers(self):
        try:
            sup = ['ALL'] + db_get_suppliers()
            self.match_supplier_combo['values'] = sup
            if not self.match_supplier_var.get() or self.match_supplier_var.get() not in sup:
                self.match_supplier_var.set('ALL')
        except Exception as e:
            print(f"[ERROR] {e}")

    def _build_tab_data(self):
        top = ttk.Frame(self.tab_data); top.pack(fill='x', pady=(0,8))
        ttk.Button(top, text="📥 Import Excel ke Database…",
                   command=self._show_import_dialog).pack(side='left', padx=4)
        ttk.Button(top, text="🔄 Refresh",
                   command=self._refresh_data_tab).pack(side='left', padx=4)
        ttk.Button(top, text="🗑 Hapus Semua Data",
                   command=self._clear_all_data).pack(side='left', padx=4)
        self.data_info_var = tk.StringVar(value="Belum ada data.")
        ttk.Label(top, textvariable=self.data_info_var, foreground='blue').pack(side='left', padx=16)

        sf = ttk.LabelFrame(self.tab_data, text="Sesi Import", padding=6)
        sf.pack(fill='x', pady=(0,8))
        self.sesi_tree = ttk.Treeview(sf, columns=('sesi','file','sheet','jml'),
                                       show='headings', height=4)
        for c,h,w in [('sesi','Sesi',160),('file','File',400),
                      ('sheet','Sheet',100),('jml','Jumlah',100)]:
            self.sesi_tree.heading(c, text=h); self.sesi_tree.column(c, width=w, anchor='w')
        self.sesi_tree.pack(fill='x')

        pf = ttk.LabelFrame(self.tab_data, text="Preview (100 baris pertama)", padding=6)
        pf.pack(fill='both', expand=True)
        cols = ('id','row_excel','tanggal','sumber','supplier','no_sumber','debit','kredit')
        heads = {'id':'ID','row_excel':'Baris','tanggal':'Tanggal','sumber':'Sumber',
                 'supplier':'Supplier','no_sumber':'No. Sumber','debit':'Debit','kredit':'Kredit'}
        widths = {'id':60,'row_excel':60,'tanggal':100,'sumber':150,'supplier':200,
                  'no_sumber':130,'debit':120,'kredit':120}
        tf = ttk.Frame(pf); tf.pack(fill='both', expand=True)
        self.preview_tree = ttk.Treeview(tf, columns=cols, show='headings', height=12)
        for c in cols:
            self.preview_tree.heading(c, text=heads[c])
            self.preview_tree.column(c, width=widths[c], anchor='w')
        self.preview_tree.column('debit', anchor='e'); self.preview_tree.column('kredit', anchor='e')
        vsb = ttk.Scrollbar(tf, orient='vertical', command=self.preview_tree.yview)
        hsb = ttk.Scrollbar(tf, orient='horizontal', command=self.preview_tree.xview)
        self.preview_tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self.preview_tree.grid(row=0,column=0,sticky='nsew')
        vsb.grid(row=0,column=1,sticky='ns'); hsb.grid(row=1,column=0,sticky='ew')
        tf.rowconfigure(0, weight=1); tf.columnconfigure(0, weight=1)

    def _refresh_data_tab(self):
        try:
            for i in self.sesi_tree.get_children(): self.sesi_tree.delete(i)
            sessions = db_get_import_sessions()
            for s in sessions:
                self.sesi_tree.insert('', 'end', values=(
                    s['sesi_import'], s['source_file'], s['sheet'], s['jml']))
            total = db_count_rows()
            self.data_info_var.set(f"Total baris: {total:,} | Sesi: {len(sessions)}")
            for i in self.preview_tree.get_children(): self.preview_tree.delete(i)
            for r in db_get_data(limit=100):
                self.preview_tree.insert('', 'end', values=(
                    r['id'], r['row_excel'], r['tanggal'], r['sumber'],
                    r['supplier'], r['no_sumber'],
                    fmt_money(r['debit']) if r['debit'] else '',
                    fmt_money(r['kredit']) if r['kredit'] else ''))
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def _build_tab_match(self):
        opt = ttk.LabelFrame(self.tab_match, text="Opsi Pencocokan", padding=10)
        opt.pack(fill='x', pady=(0,8))

        f1 = ttk.Frame(opt); f1.pack(fill='x', pady=4)
        ttk.Label(f1, text="Supplier:", width=16).pack(side='left')
        self.match_supplier_var = tk.StringVar(value='ALL')
        self.match_supplier_combo = ttk.Combobox(f1, textvariable=self.match_supplier_var,
                                                   state='readonly', width=60)
        self.match_supplier_combo.pack(side='left', padx=4)
        ttk.Button(f1, text="🔄", width=3,
                   command=self._refresh_suppliers).pack(side='left', padx=2)

        f2 = ttk.Frame(opt); f2.pack(fill='x', pady=4)
        ttk.Label(f2, text="Maks. baris per grup:", width=16).pack(side='left')
        self.max_group_var = tk.StringVar(value='30')
        ttk.Combobox(f2, textvariable=self.max_group_var, state='readonly',
                     width=5, values=['1','2','3','4','5','6','7','8','9','10',
                                      '12','15','20','25','30','40','50','75','100']
                     ).pack(side='left', padx=4)

        f3 = ttk.Frame(opt); f3.pack(fill='x', pady=4)
        ttk.Label(f3, text="Time budget (ms/debit):", width=16).pack(side='left')
        self.time_budget_var = tk.StringVar(value='2000')
        ttk.Combobox(f3, textvariable=self.time_budget_var, state='readonly',
                     width=8, values=['50','100','200','500','1000','2000','5000','10000']
                     ).pack(side='left', padx=4)
        ttk.Label(f3, text="(kecil=cepat tapi mungkin lewat, besar=lambat tapi lengkap)",
                  foreground='gray').pack(side='left', padx=6)

        f4 = ttk.Frame(opt); f4.pack(fill='x', pady=(10,0))
        self.match_progress = ttk.Progressbar(f4, mode='determinate', length=400)
        self.match_progress.pack(side='left', fill='x', expand=True, padx=(0,8))
        self.match_progress_label = ttk.Label(f4, text="", foreground='gray', width=55)
        self.match_progress_label.pack(side='left')

        f5 = ttk.Frame(opt); f5.pack(fill='x', pady=(10,0))
        self.btn_run_match = ttk.Button(f5, text="▶  Jalankan Pencocokan",
                                         command=self._run_matching)
        self.btn_run_match.pack(side='right', padx=4)
        self.btn_cancel_match = ttk.Button(f5, text="⏹ Cancel",
                                            command=self._cancel_matching,
                                            state='disabled')
        self.btn_cancel_match.pack(side='right', padx=4)

        info = ttk.LabelFrame(self.tab_match, text="Ringkasan Terakhir", padding=10)
        info.pack(fill='both', expand=True)
        self.match_summary_var = tk.StringVar(value="Belum ada pencocokan.")
        ttk.Label(info, textvariable=self.match_summary_var,
                  font=('Segoe UI', 10), justify='left').pack(anchor='w')

    def _build_tab_result(self):
        top = ttk.Frame(self.tab_result); top.pack(fill='x', pady=(0,8))
        ttk.Label(top, text="Tampilkan:").pack(side='left')
        self.result_filter_var = tk.StringVar(value='Semua')
        fc = ttk.Combobox(top, textvariable=self.result_filter_var, state='readonly',
                          width=20, values=['Semua','Cocok','Tanpa Pasangan'])
        fc.pack(side='left', padx=6)
        fc.bind('<<ComboboxSelected>>', lambda e: self._render_results())
        ttk.Button(top, text="🎨 Ekspor Excel Bertanda",
                   command=self._export_marked).pack(side='right', padx=4)
        ttk.Button(top, text="📄 Ekspor Excel (ringkas)",
                   command=self._export_ringkas).pack(side='right', padx=4)
        ttk.Button(top, text="📋 Ekspor CSV",
                   command=self._export_csv).pack(side='right', padx=4)
        self.result_summary_var = tk.StringVar(value="")
        ttk.Label(self.tab_result, textvariable=self.result_summary_var,
                  font=('Segoe UI', 9, 'bold')).pack(anchor='w', pady=(0,6))

        cols = ('status','supplier','tipe','nilai','baris_debit','baris_kredit')
        heads = {'status':'Status','supplier':'Supplier','tipe':'Tipe','nilai':'Nilai',
                 'baris_debit':'Baris Debit','baris_kredit':'Baris Kredit'}
        widths = {'status':140,'supplier':220,'tipe':80,'nilai':140,
                  'baris_debit':200,'baris_kredit':200}
        tf = ttk.Frame(self.tab_result); tf.pack(fill='both', expand=True)
        self.result_tree = ttk.Treeview(tf, columns=cols, show='headings', height=18)
        for c in cols:
            self.result_tree.heading(c, text=heads[c])
            self.result_tree.column(c, width=widths[c], anchor='w')
        self.result_tree.column('nilai', anchor='e')
        vsb = ttk.Scrollbar(tf, orient='vertical', command=self.result_tree.yview)
        hsb = ttk.Scrollbar(tf, orient='horizontal', command=self.result_tree.xview)
        self.result_tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self.result_tree.grid(row=0,column=0,sticky='nsew')
        vsb.grid(row=0,column=1,sticky='ns'); hsb.grid(row=1,column=0,sticky='ew')
        tf.rowconfigure(0, weight=1); tf.columnconfigure(0, weight=1)
        self.result_tree.tag_configure('match', background='#e7f6e7')
        self.result_tree.tag_configure('unmatch', background='#ffe0e0')

    def _render_results(self):
        for i in self.result_tree.get_children(): self.result_tree.delete(i)
        if not self.last_match_sesi: return
        fv = self.result_filter_var.get()
        conn = db_connect()
        if fv in ('Semua', 'Cocok'):
            for r in conn.execute("""SELECT supplier, tipe_group, total_nilai,
                                            debit_ids, credit_ids FROM hasil_match
                                     WHERE sesi_match = ? ORDER BY id""", (self.last_match_sesi,)):
                d_rows = self._get_rows(r['debit_ids']); c_rows = self._get_rows(r['credit_ids'])
                self.result_tree.insert('', 'end', values=(
                    'COCOK', r['supplier'], r['tipe_group'], fmt_money(r['total_nilai']),
                    ', '.join(str(x) for x in d_rows), ', '.join(str(x) for x in c_rows)),
                    tags=('match',))
        if fv in ('Semua', 'Tanpa Pasangan'):
            for r in conn.execute("""SELECT supplier, row_excel, nilai, sisi FROM tanpa_pasangan
                                     WHERE sesi_match = ? ORDER BY id""", (self.last_match_sesi,)):
                d = str(r['row_excel']) if r['sisi'] == 'debit' else ''
                c = str(r['row_excel']) if r['sisi'] == 'kredit' else ''
                self.result_tree.insert('', 'end', values=(
                    'TANPA PASANGAN', r['supplier'], r['sisi'], fmt_money(r['nilai']), d, c),
                    tags=('unmatch',))
        conn.close()

    def _get_rows(self, ids_str):
        if not ids_str: return []
        ids = [int(x) for x in ids_str.split(',') if x.strip()]
        if not ids: return []
        conn = db_connect()
        ph = ','.join('?'*len(ids))
        m = {r[0]: r[1] for r in conn.execute(
            f"SELECT id, row_excel FROM transaksi WHERE id IN ({ph})", ids)}
        conn.close()
        return [m.get(i, '?') for i in ids]

    def _show_import_dialog(self):
        ImportDialog(self)

    def _clear_all_data(self):
        if not messagebox.askyesno("Konfirmasi", "Yakin hapus SEMUA data?"): return
        db_clear_all()
        self.last_match_sesi = None
        self.match_summary_var.set("Data dihapus.")
        self._refresh_data_tab(); self._refresh_suppliers()
        self.status_var.set("Database dikosongkan.")

    def _run_matching(self):
        if self.worker_thread and self.worker_thread.is_alive():
            self.cancel_event.set()
            self.status_var.set("Menghentikan proses lama...")
            self.update()
            self.worker_thread.join(timeout=3.0)
            if self.worker_thread.is_alive():
                messagebox.showwarning("Perhatian",
                    "Proses lama belum berhenti. Tunggu beberapa detik.")
                return

        supplier = self.match_supplier_var.get()
        try: max_group = int(self.max_group_var.get())
        except ValueError: max_group = 30
        try: time_budget = int(self.time_budget_var.get()) / 1000.0
        except ValueError: time_budget = 2.0

        try:
            rows = db_get_data(supplier=supplier if supplier != 'ALL' else None)
        except Exception as e:
            messagebox.showerror("Error", str(e)); return
        if not rows:
            messagebox.showwarning("Perhatian", "Tidak ada data."); return

        self.cancel_event.clear()
        self.btn_run_match.configure(state='disabled')
        self.btn_cancel_match.configure(state='normal')
        self.match_progress['value'] = 0
        self.match_progress_label.configure(text=f"Memproses {len(rows)} baris...")
        self.status_var.set(f"Memproses {len(rows)} baris...")

        cancel_evt = self.cancel_event; ui_q = self.ui_queue

        def worker():
            try:
                def cb(p, m): ui_q.put(('match_progress', (p, m)))
                g, ud, uc = match_rows(rows, max_group=max_group,
                                        time_budget=time_budget,
                                        cancel_event=cancel_evt, progress_cb=cb)
                ui_q.put(('match_done', (g, ud, uc)))
            except InterruptedError:
                ui_q.put(('match_error', 'CANCELLED'))
            except Exception as e:
                import traceback
                ui_q.put(('match_error', f"{e}\n\n{traceback.format_exc()}"))

        self.worker_thread = threading.Thread(target=worker, daemon=True)
        self.worker_thread.start()

    def _cancel_matching(self):
        if self.worker_thread and self.worker_thread.is_alive():
            self.cancel_event.set()
            self.status_var.set("Membatalkan...")

    def _save_match(self, groups, ud, uc):
        sesi = datetime.now().strftime("%Y%m%d_%H%M%S")
        conn = db_connect(); cur = conn.cursor()
        for g in groups:
            cur.execute("""INSERT INTO hasil_match
                (sesi_match, supplier, tipe_group, debit_ids, credit_ids, total_nilai)
                VALUES (?,?,?,?,?,?)""", (sesi, g['supplier'], g['size'],
                ','.join(str(x) for x in g['debit_ids']),
                ','.join(str(x) for x in g['credit_ids']), g['total']))
        for d in ud:
            cur.execute("""INSERT INTO tanpa_pasangan
                (sesi_match, supplier, transaksi_id, row_excel, nilai, sisi)
                VALUES (?,?,?,?,?,?)""", (sesi, d['raw']['supplier'], d['id'],
                d['row'], d['cents']/100.0, 'debit'))
        for c in uc:
            cur.execute("""INSERT INTO tanpa_pasangan
                (sesi_match, supplier, transaksi_id, row_excel, nilai, sisi)
                VALUES (?,?,?,?,?,?)""", (sesi, c['raw']['supplier'], c['id'],
                c['row'], c['cents']/100.0, 'kredit'))
        conn.commit(); conn.close()
        return sesi

    def _on_match_done(self, payload):
        groups, ud, uc = payload
        try: sesi = self._save_match(groups, ud, uc); self.last_match_sesi = sesi
        except Exception as e:
            messagebox.showerror("Error", str(e)); return
        n_dr = sum(len(g['debit_ids']) for g in groups)
        n_cr = sum(len(g['credit_ids']) for g in groups)
        by_type = defaultdict(int)
        for g in groups: by_type[g['size']] += 1
        lines = [f"Sesi: {sesi}", f"Total grup cocok: {len(groups)}",
                 f"  Debit masuk grup: {n_dr}", f"  Kredit masuk grup: {n_cr}",
                 f"Tanpa pasangan Debit: {len(ud)}", f"Tanpa pasangan Kredit: {len(uc)}",
                 "", "Rincian:"]
        for t, c in sorted(by_type.items()): lines.append(f"  {t}: {c} grup")
        self.match_summary_var.set('\n'.join(lines))
        self.match_progress['value'] = 100
        self.match_progress_label.configure(text="Selesai")
        self.btn_run_match.configure(state='normal')
        self.btn_cancel_match.configure(state='disabled')
        self.status_var.set(f"Selesai — {len(groups)} grup cocok.")
        self._render_results(); self.nb.select(self.tab_result)

    def _on_match_error(self, e):
        self.btn_run_match.configure(state='normal')
        self.btn_cancel_match.configure(state='disabled')
        self.match_progress['value'] = 0
        if e == 'CANCELLED':
            self.status_var.set("Dibatalkan.")
            self.match_summary_var.set("Dibatalkan.")
        else:
            messagebox.showerror("Error", e)

    def _poll_queue(self):
        try:
            while True:
                kind, payload = self.ui_queue.get_nowait()
                if kind == 'import_progress':
                    pct, msg = payload
                    try:
                        if hasattr(self, '_import_progress_bar'):
                            self._import_progress_bar['value'] = pct
                            self._import_progress_label.configure(text=msg[:60])
                    except tk.TclError: pass
                elif kind == 'import_done': self._on_import_done(payload)
                elif kind == 'import_error': self._on_import_error(payload)
                elif kind == 'match_progress':
                    pct, msg = payload
                    self.match_progress['value'] = pct
                    self.match_progress_label.configure(text=msg[:60])
                elif kind == 'match_done': self._on_match_done(payload)
                elif kind == 'match_error': self._on_match_error(payload)
        except queue.Empty: pass
        self.after(80, self._poll_queue)

    def _on_import_done(self, result):
        for w in self.winfo_children():
            if isinstance(w, ImportDialog):
                try: w.destroy()
                except Exception: pass
        self._refresh_data_tab(); self._refresh_suppliers()
        messagebox.showinfo("Import Sukses",
            f"Sesi: {result['sesi']}\nBaris ditambahkan: {result['inserted']:,}\n"
            f"Baris dilewati: {result['skipped']:,}\nTotal dibaca: {result['total_rows']:,}")

    def _on_import_error(self, e):
        for w in self.winfo_children():
            if isinstance(w, ImportDialog):
                try: w.destroy()
                except Exception: pass
        messagebox.showerror("Import Gagal", e)

    def _get_results_data(self):
        if not self.last_match_sesi: return None
        conn = db_connect()
        matches = []
        for r in conn.execute("""SELECT supplier, tipe_group, total_nilai,
                                        debit_ids, credit_ids FROM hasil_match
                                 WHERE sesi_match = ? ORDER BY id""", (self.last_match_sesi,)):
            matches.append({'supplier': r['supplier'], 'tipe': r['tipe_group'],
                            'nilai': r['total_nilai'],
                            'baris_debit': self._get_rows(r['debit_ids']),
                            'baris_kredit': self._get_rows(r['credit_ids'])})
        unmatched = [{'supplier': r['supplier'], 'row': r['row_excel'],
                      'nilai': r['nilai'], 'sisi': r['sisi']}
                     for r in conn.execute("""SELECT supplier, row_excel, nilai, sisi
                                              FROM tanpa_pasangan WHERE sesi_match = ?
                                              ORDER BY id""", (self.last_match_sesi,))]
        conn.close()
        return {'matches': matches, 'unmatched': unmatched}

    def _export_ringkas(self):
        data = self._get_results_data()
        if not data:
            messagebox.showinfo("Info", "Belum ada hasil."); return
        if not HAS_OPENPYXL:
            messagebox.showerror("Error", "openpyxl tidak ada."); return
        path = filedialog.asksaveasfilename(title="Simpan ringkas", defaultextension='.xlsx',
            filetypes=[("Excel","*.xlsx")], initialfile=f"ringkas_{self.last_match_sesi}.xlsx")
        if not path: return
        try:
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Hasil"
            hd = ['Status','Supplier','Tipe','Nilai','Baris Debit','Baris Kredit']
            ws.append(hd)
            hf = Font(bold=True, color='FFFFFF'); hfill = PatternFill('solid', fgColor='305496')
            for c in range(1, len(hd)+1):
                cell = ws.cell(row=1, column=c); cell.font = hf; cell.fill = hfill
            fm = PatternFill('solid', fgColor='D9EAD3')
            fu = PatternFill('solid', fgColor='F4CCCC')
            for m in data['matches']:
                ws.append(['COCOK', m['supplier'], m['tipe'], fmt_money(m['nilai']),
                           ', '.join(str(x) for x in m['baris_debit']),
                           ', '.join(str(x) for x in m['baris_kredit'])])
                for c in range(1, 7): ws.cell(row=ws.max_row, column=c).fill = fm
            for u in data['unmatched']:
                d = str(u['row']) if u['sisi'] == 'debit' else ''
                c_ = str(u['row']) if u['sisi'] == 'kredit' else ''
                ws.append(['TANPA PASANGAN', u['supplier'], u['sisi'],
                           fmt_money(u['nilai']), d, c_])
                for c in range(1, 7): ws.cell(row=ws.max_row, column=c).fill = fu
            for i, w in enumerate([18,30,10,16,22,22], start=1):
                ws.column_dimensions[get_column_letter(i)].width = w
            wb.save(path)
            messagebox.showinfo("Sukses", f"Disimpan:\n{path}")
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def _export_csv(self):
        data = self._get_results_data()
        if not data:
            messagebox.showinfo("Info", "Belum ada hasil."); return
        path = filedialog.asksaveasfilename(title="Simpan CSV", defaultextension='.csv',
            filetypes=[("CSV","*.csv")], initialfile=f"hasil_{self.last_match_sesi}.csv")
        if not path: return
        try:
            with open(path, 'w', encoding='utf-8-sig', newline='') as f:
                w = csv.writer(f, delimiter=';')
                w.writerow(['Status','Supplier','Tipe','Nilai','Baris Debit','Baris Kredit'])
                for m in data['matches']:
                    w.writerow(['COCOK', m['supplier'], m['tipe'], fmt_money(m['nilai']),
                                ', '.join(str(x) for x in m['baris_debit']),
                                ', '.join(str(x) for x in m['baris_kredit'])])
                for u in data['unmatched']:
                    d = str(u['row']) if u['sisi'] == 'debit' else ''
                    c_ = str(u['row']) if u['sisi'] == 'kredit' else ''
                    w.writerow(['TANPA PASANGAN', u['supplier'], u['sisi'],
                                fmt_money(u['nilai']), d, c_])
            messagebox.showinfo("Sukses", f"Disimpan:\n{path}")
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def _export_marked(self):
        if not self.imported_file or not os.path.exists(self.imported_file):
            messagebox.showwarning("Perhatian", "File Excel sumber tidak diketahui."); return
        if not HAS_OPENPYXL:
            messagebox.showerror("Error", "openpyxl tidak ada."); return
        if not self.last_match_sesi:
            messagebox.showinfo("Info", "Belum ada hasil."); return
        ext = os.path.splitext(self.imported_file)[1].lower()
        if ext not in ('.xlsx', '.xlsm'):
            messagebox.showwarning("Perhatian", "Hanya untuk .xlsx / .xlsm."); return
        base = os.path.splitext(os.path.basename(self.imported_file))[0]
        path = filedialog.asksaveasfilename(title="Simpan bertanda", defaultextension=ext,
            filetypes=[("Excel", f"*{ext}")],
            initialfile=f"{base}_BERTANDA_{datetime.now():%Y%m%d_%H%M%S}{ext}")
        if not path: return
        try:
            conn = db_connect(); cur = conn.cursor()
            row_status = {}
            for m in cur.execute("""SELECT tipe_group, total_nilai, debit_ids, credit_ids
                                     FROM hasil_match WHERE sesi_match = ?""",
                                  (self.last_match_sesi,)):
                d_ids = [int(x) for x in m['debit_ids'].split(',') if x.strip()]
                c_ids = [int(x) for x in m['credit_ids'].split(',') if x.strip()]
                d_rows = self._get_rows(m['debit_ids']); c_rows = self._get_rows(m['credit_ids'])
                for did in d_ids:
                    r = cur.execute("SELECT row_excel FROM transaksi WHERE id=?", (did,)).fetchone()
                    if r: row_status[r[0]] = ('COCOK', ', '.join(str(x) for x in c_rows),
                                              fmt_money(m['total_nilai']))
                for cid in c_ids:
                    r = cur.execute("SELECT row_excel FROM transaksi WHERE id=?", (cid,)).fetchone()
                    if r: row_status[r[0]] = ('COCOK', ', '.join(str(x) for x in d_rows),
                                              fmt_money(m['total_nilai']))
            for r in cur.execute("SELECT row_excel FROM tanpa_pasangan WHERE sesi_match=?",
                                  (self.last_match_sesi,)):
                row_status[r[0]] = ('TANPA PASANGAN', '', '')
            conn.close()

            wb = openpyxl.load_workbook(self.imported_file, keep_vba=(ext=='.xlsm'))
            sheet_name = self.imported_sheet
            ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
            hr = 1
            last_col = ws.max_column
            cs = last_col+1; cp = last_col+2; ca = last_col+3
            hf = Font(bold=True, color='FFFFFF'); hfill = PatternFill('solid', fgColor='305496')
            for c, name in [(cs,'Status'),(cp,'Baris Pasangan'),(ca,'Nilai Pasangan')]:
                cell = ws.cell(row=hr, column=c, value=name)
                cell.font = hf; cell.fill = hfill
            fm = PatternFill('solid', fgColor='C6EFCE')
            fu = PatternFill('solid', fgColor='FFC7CE')
            fontm = Font(color='006100', bold=True); fontu = Font(color='9C0006', bold=True)
            for row_idx in range(hr+1, ws.max_row+1):
                info = row_status.get(row_idx)
                if not info: continue
                status, pr, pa = info
                ws.cell(row=row_idx, column=cs, value=status)
                ws.cell(row=row_idx, column=cp, value=pr)
                ws.cell(row=row_idx, column=ca, value=pa)
                f = fm if status == 'COCOK' else fu
                fo = fontm if status == 'COCOK' else fontu
                for c in range(1, ca+1):
                    cell = ws.cell(row=row_idx, column=c); cell.fill = f; cell.font = fo
            ws.column_dimensions[get_column_letter(cs)].width = 16
            ws.column_dimensions[get_column_letter(cp)].width = 18
            ws.column_dimensions[get_column_letter(ca)].width = 18
            wb.save(path)
            messagebox.showinfo("Sukses", f"Disimpan:\n{path}")
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def _show_about(self):
        scipy_txt = "Aktif ✓" if HAS_SCIPY else "Tidak aktif (install: pip install scipy)"
        messagebox.showinfo("Tentang",
            "Pencocokan Transaksi v4.3.2\n"
            "DFS + MILP + Deep Search\n\n"
            f"Status MILP: {scipy_txt}\n\n"
            "Perbaikan v4.3.2:\n"
            "• MILP timeout diperpanjang (maks 60 detik)\n"
            "• PASS 7 Deep Search: MILP 30 detik untuk sisa debit\n"
            "• Progress per debit di setiap pass\n"
            "• HiGHS output di-suppress\n"
            f"\nDatabase: {DB_FILE}")


# =========================================================
# DIALOG IMPORT
# =========================================================
class ImportDialog(tk.Toplevel):
    def __init__(self, parent):
        super().__init__(parent)
        self.title("Import Excel ke Database")
        self.geometry("720x560")
        self.transient(parent); self.grab_set()
        self.parent = parent
        self.raw_rows = []; self.headers = []; self.sheet_names = []

        f = ttk.Frame(self, padding=10); f.pack(fill='both', expand=True)

        row1 = ttk.Frame(f); row1.pack(fill='x', pady=3)
        ttk.Label(row1, text="File Excel:", width=14).pack(side='left')
        self.path_var = tk.StringVar()
        ttk.Entry(row1, textvariable=self.path_var).pack(side='left', fill='x', expand=True, padx=2)
        ttk.Button(row1, text="Pilih…", command=self.browse).pack(side='left')

        row2 = ttk.Frame(f); row2.pack(fill='x', pady=3)
        ttk.Label(row2, text="Sheet:", width=14).pack(side='left')
        self.sheet_var = tk.StringVar()
        self.sheet_combo = ttk.Combobox(row2, textvariable=self.sheet_var, state='readonly')
        self.sheet_combo.pack(side='left', fill='x', expand=True, padx=2)
        self.sheet_combo.bind('<<ComboboxSelected>>', lambda e: self.load_file())

        row3 = ttk.Frame(f); row3.pack(fill='x', pady=3)
        ttk.Label(row3, text="Baris Header:", width=14).pack(side='left')
        self.header_row_var = tk.StringVar(value='1')
        e = ttk.Entry(row3, textvariable=self.header_row_var, width=6); e.pack(side='left', padx=2)
        e.bind('<Return>', lambda ev: self.load_file())
        ttk.Button(row3, text="Muat", command=self.load_file, width=8).pack(side='left', padx=2)
        self.info_var = tk.StringVar(value="Belum ada file.")
        ttk.Label(row3, textvariable=self.info_var, foreground='gray').pack(side='left', padx=6)

        mf = ttk.LabelFrame(f, text="Pemetaan Kolom", padding=8); mf.pack(fill='x', pady=8)
        ttk.Label(mf, text="Supplier:").grid(row=0,column=0,sticky='w',pady=2)
        self.col_supplier = ttk.Combobox(mf, state='readonly'); self.col_supplier.grid(row=0,column=1,sticky='ew',padx=4)
        ttk.Label(mf, text="Debit:").grid(row=1,column=0,sticky='w',pady=2)
        self.col_debit = ttk.Combobox(mf, state='readonly'); self.col_debit.grid(row=1,column=1,sticky='ew',padx=4)
        ttk.Label(mf, text="Kredit:").grid(row=2,column=0,sticky='w',pady=2)
        self.col_credit = ttk.Combobox(mf, state='readonly'); self.col_credit.grid(row=2,column=1,sticky='ew',padx=4)
        mf.columnconfigure(1, weight=1)

        opt = ttk.Frame(f); opt.pack(fill='x', pady=6)
        self.replace_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(opt, text="Hapus data lama sebelum import",
                        variable=self.replace_var).pack(anchor='w')

        pf = ttk.Frame(f); pf.pack(fill='x', pady=6)
        self.parent._import_progress_bar = ttk.Progressbar(pf, mode='determinate', length=400)
        self.parent._import_progress_bar.pack(side='left', fill='x', expand=True, padx=(0,6))
        self.parent._import_progress_label = ttk.Label(pf, text="", foreground='gray')
        self.parent._import_progress_label.pack(side='left')

        bf = ttk.Frame(f); bf.pack(fill='x', pady=8)
        ttk.Button(bf, text="▶  Import Sekarang", command=self.do_import).pack(side='right', padx=4)
        ttk.Button(bf, text="Batal", command=self.destroy).pack(side='right', padx=4)

    def browse(self):
        p = filedialog.askopenfilename(title="Pilih file Excel",
            filetypes=[("Excel","*.xlsx *.xlsm"),("Semua","*.*")])
        if p: self.path_var.set(p); self.load_file()

    def load_file(self):
        path = self.path_var.get()
        if not path or not os.path.exists(path): return
        if not HAS_OPENPYXL: return
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            self.sheet_names = wb.sheetnames; wb.close()
        except Exception as e:
            messagebox.showerror("Error", f"Gagal buka file:\n{e}"); return
        self.sheet_combo['values'] = self.sheet_names
        if self.sheet_names and not self.sheet_var.get():
            self.sheet_var.set(self.sheet_names[0])
        try:
            self.raw_rows = read_excel_rows(path, self.sheet_var.get())
        except Exception as e:
            messagebox.showerror("Error", f"Gagal baca:\n{e}"); return
        try: hr = max(1, int(self.header_row_var.get()))
        except ValueError: hr = 1
        if self.raw_rows and hr <= len(self.raw_rows):
            self.headers = [str(h).strip() if h is not None and str(h).strip() else f"Kolom{i+1}"
                            for i, h in enumerate(self.raw_rows[hr-1])]
        else: self.headers = []
        self.info_var.set(f"{len(self.raw_rows)-hr} baris, {len(self.headers)} kolom")
        for cb in (self.col_supplier, self.col_debit, self.col_credit):
            cb['values'] = self.headers
        def find(*keys):
            for k in keys:
                for h in self.headers:
                    if k.lower() in h.lower(): return h
            return ''
        if not self.col_supplier.get():
            v = find('supplier','vendor','pemasok','pelanggan','nama pelanggan')
            if v: self.col_supplier.set(v)
        if not self.col_debit.get():
            v = find('debit')
            if v: self.col_debit.set(v)
        if not self.col_credit.get():
            v = find('kredit')
            if v: self.col_credit.set(v)

    def do_import(self):
        path = self.path_var.get()
        if not path or not os.path.exists(path):
            messagebox.showwarning("Perhatian", "Pilih file."); return
        if not self.col_supplier.get():
            messagebox.showwarning("Perhatian", "Kolom Supplier wajib."); return
        if not (self.col_debit.get() or self.col_credit.get()):
            messagebox.showwarning("Perhatian", "Pilih Debit/Kredit."); return
        col_map = {'supplier': self.col_supplier.get(),
                   'debit': self.col_debit.get(), 'kredit': self.col_credit.get()}
        self.parent.imported_file = path
        self.parent.imported_sheet = self.sheet_var.get()
        replace = self.replace_var.get()
        header_row = self.header_row_var.get()
        sheet = self.sheet_var.get()
        label = os.path.basename(path)
        self.parent.cancel_event.clear()

        def worker():
            try:
                def cb(p, m): self.parent.ui_queue.put(('import_progress', (p, m)))
                result = db_import_excel(path, sheet, header_row, col_map, label,
                                          replace=replace, progress_cb=cb,
                                          cancel_event=self.parent.cancel_event)
                self.parent.ui_queue.put(('import_done', result))
            except InterruptedError:
                self.parent.ui_queue.put(('import_error', 'Import dibatalkan.'))
            except Exception as e:
                self.parent.ui_queue.put(('import_error', str(e)))

        threading.Thread(target=worker, daemon=True).start()


def main():
    if not HAS_OPENPYXL:
        print("[PERINGATAN] openpyxl tidak ada.")
    if not HAS_SCIPY:
        print("[INFO] scipy tidak terdeteksi. Kombinasi besar mungkin tidak ketemu.")
        print("       Install dengan: pip install scipy")
    App().mainloop()


if __name__ == '__main__':
    main()